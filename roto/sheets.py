"""Upload an xlsx workbook to Google Drive as a native Google Sheet.

Uses the installed-app OAuth client from ~/Downloads (project: hootro) with
the drive.file scope (can only touch files it creates). First run opens a
browser for consent; the token is cached at ~/.config/cube-roto/token.json.

Usage: python3 sheets.py workbook.xlsx "Sheet Title"
"""

import glob
import json
import pathlib
import sys

from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from googleapiclient.http import MediaFileUpload

SCOPES = ["https://www.googleapis.com/auth/drive.file"]
TOKEN_PATH = pathlib.Path.home() / ".config" / "cube-roto" / "token.json"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEET_MIME = "application/vnd.google-apps.spreadsheet"


def find_client_secret():
    for f in sorted(glob.glob(str(pathlib.Path.home() / "Downloads" / "client_secret*.json"))):
        if "installed" in json.loads(pathlib.Path(f).read_text()):
            return f
    sys.exit("no installed-app client_secret*.json found in ~/Downloads")


def get_credentials():
    creds = None
    if TOKEN_PATH.exists():
        creds = Credentials.from_authorized_user_file(TOKEN_PATH, SCOPES)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            flow = InstalledAppFlow.from_client_secrets_file(find_client_secret(), SCOPES)
            creds = flow.run_local_server(port=0)
        TOKEN_PATH.parent.mkdir(parents=True, exist_ok=True)
        TOKEN_PATH.write_text(creds.to_json())
    return creds


def upload(path, title, file_id=None):
    """Create a new Sheet, or replace the content of an existing one in place."""
    drive = build("drive", "v3", credentials=get_credentials())
    media = MediaFileUpload(path, mimetype=XLSX_MIME)
    if file_id:
        req = drive.files().update(fileId=file_id, media_body=media, fields="id,webViewLink")
    else:
        req = drive.files().create(
            body={"name": title, "mimeType": SHEET_MIME}, media_body=media,
            fields="id,webViewLink",
        )
    file = req.execute()
    print(file["webViewLink"])
    return file


def verify(formulas_path, values_path, *, tabs, fuzzy_headers):
    """Upload formula build to a temp Sheet, export Google's evaluated
    values, and diff against the values build cell by cell."""
    import io

    import openpyxl
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaIoBaseDownload

    drive = build("drive", "v3", credentials=get_credentials())
    tmp = upload(str(formulas_path), "tmp-roto-verify")
    try:
        req = drive.files().export_media(
            fileId=tmp["id"],
            mimeType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        buf = io.BytesIO()
        dl = MediaIoBaseDownload(buf, req)
        done = False
        while not done:
            _, done = dl.next_chunk()
    finally:
        drive.files().delete(fileId=tmp["id"]).execute()

    wb_want = openpyxl.load_workbook(values_path, data_only=True)
    wb_got = openpyxl.load_workbook(io.BytesIO(buf.getvalue()), data_only=True)
    total, mismatches = 0, []
    for tab in tabs:
        want, got = wb_want[tab], wb_got[tab]
        assert want.max_row == got.max_row, (tab, want.max_row, got.max_row)
        headers = [c.value or "" for c in want[1]]
        fuzzy = {i + 1 for i, h in enumerate(headers)
                 if any(k in str(h) for k in fuzzy_headers)}
        for r in range(1, want.max_row + 1):
            for c in range(1, len(headers) + 1):
                total += 1
                w, g = want.cell(r, c).value, got.cell(r, c).value
                if g == "":
                    g = None
                if w == g:
                    continue
                if (c in fuzzy and isinstance(w, (int, float))
                        and isinstance(g, (int, float)) and abs(w - g) <= 0.051):
                    continue
                mismatches.append((tab, r, c, w, g))
    if mismatches:
        for m in mismatches[:20]:
            print("MISMATCH %s row %d col %d: values=%r evaluated=%r" % m)
        sys.exit(f"verification FAILED: {len(mismatches)}/{total} cells differ")
    print(f"verification passed: {total} cells match")



def main():
    args = sys.argv[1:]
    file_id = None
    if "--update" in args:
        i = args.index("--update")
        file_id = args[i + 1]
        del args[i : i + 2]
    if len(args) != 2:
        sys.exit(__doc__)
    upload(args[0], args[1], file_id)


if __name__ == "__main__":
    main()
