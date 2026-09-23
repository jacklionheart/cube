"""Refresh the ALL-DRAFTS Samp Cube Roto sheet (every roto, all seasons).

The default 13-pod s4 sheet is refresh.py's job and is never touched
here. This builds a second workbook over every complete draft known to
read-the-bones or the Google Sheets (37 drafts, Oct 2025 - Sep 2026)
and publishes it in place to its own Sheet.

Per-draft source and double-pick round (dpa) are pinned below.
RTB-sourced drafts with a redacted seat use the full Google Sheet when
one exists; three redacted drafts (Birds of Paradise, Ravnica,
Innistrad: Dark Ascension) have no sheet, so one player's picks are
blank there. Teams tabs scale to ~1/3 of decked drafts (--teams-frac).

Usage: python3 refresh_all.py [--verify] [--dry-run]
"""

import json
import pathlib
import subprocess
import sys

import refresh
from refresh import COMPUTED_TABS, FUZZY_HEADERS, curl, slugify, verify

HERE = pathlib.Path(__file__).parent

# name -> (source, dpa). source: "rtb:<slug>" or a Google Sheet id.
# dpa = double-pick-after round from RTB's declaration (validated
# against pickN); None would mean plain snake.
SOURCES_ALL = {
    "Lightning Bolt (s1)": ("rtb:lightning-bolt", 20),
    "Birds of Paradise (s1)": ("rtb:birds-of-paradise", 20),
    "Zendikar (s2)": ("rtb:zendikar", 25),
    "Tarkir (s2)": ("rtb:tarkir", 25),
    "Ravnica (s2)": ("rtb:ravnica", 25),
    "Lorwyn (s2)": ("rtb:lorwyn", 25),
    "Innistrad (s2)": ("1aYwiIZhjAs4y3rwQPrTNBwKyzLV9P5GUny2MuDm29xI", 25),
    "Pyrogoyf (s2)": ("1U6xwHq7etNRtmcU5ez8YiCOgcoJQaOhG10i7k9Jn3MU", 25),
    "Dark Ascension (s2)": ("rtb:innistrad-dark-ascension", 25),
    "Fate Reforged (s2)": ("rtb:tarkir-fate-reforged", 25),
    "Bloodbraid Elf (s3)": ("1yPeuxcrp6ELYZhYe22CRCvMnYwVGAXVjjNhkGsSH7Nk", 25),
    "Thoughtseize (s3)": ("10RYiEclXvYNIyBX5Umhwr6ZVjVfYyZR6Hk0ytNMyGNk", 25),
    "Tarmogoyf (s3)": ("1_Cs54SgPgqwazqCUMOrJwK6uppnVYAKMtZEWUNHmwuo", 25),
    "Dark Confidant (s3)": ("19oKmIFp2SvrXvhbBYNn3XCqywhF6qf-wSDQu1hG-BpU", 25),
    "Blightning (s3)": ("1ukMEGa5uzweV11BY0QvZKb7VZ6DP6YOQbKU69h2Iwfs", 25),
    "Terminate (s3)": ("1vVTTwOzFO1KMZz8kOIQC-wKOk0ZtmNj1hJY4NgqWDTY", 24),
    "Maelstrom Pulse (s3)": ("1s-TWSw3TpGSBnrRxPto0lmQ5cVYs4u5pT4jdFyNPO0U", 25),
    "Liliana (s3)": ("1sJo9HL0M4b-moD4Cs1_SkMDcc1FtNccEH6rsgukVShs", 25),
    "Raging Ravine (s3)": ("1tV7db2X21V1uWwRDe_Z4aBLNMJWMbjbXKlIrE2gy_ZE", 25),
    "Scavenging Ooze (s3)": ("rtb:scavenging-ooze", 20),
    "Cabal (s3)": ("1OSaWnKMEbaCXIdhqKxBKQ8Iy99fOsz7otLmTXPZDLk8", 25),
    "Inquisition (s3)": ("1N_3n6rH-FsJAeXj44KjLGcwmXc7XIkY0_yPy437FkFw", 25),
    "Rabbit Battery (s3)": ("1RXPYzorE-nN9UCpvp0KG-KqP5VNBjSH80tTmtdFGKF0", 25),
    "Kappa Cannoneer (s3)": ("rtb:kappa-cannoneer", 23),
}
# the 13 s4 pods keep their exact refresh.py names so decks.tsv matches
SOURCES_ALL.update({n: (sid, 25) for n, sid in refresh.SOURCES.items()})

TEAMS_FRAC = 0.33
# Created once via upload_sheet.py; update in place forever after.
TARGET_SHEET_ID = "1eG04gK3iI3jlLDTSrlArgDEkEmmEHrSW4Hwq-d1SzHI"
TITLE = "Samp Cube Roto — All Drafts"


def download_sources(src_dir):
    src_dir.mkdir(exist_ok=True)
    for name, (sid, dpa) in SOURCES_ALL.items():
        out = src_dir / f"{slugify(name)}.xlsx"
        if sid.startswith("rtb:"):
            subprocess.run(
                [sys.executable, str(HERE / "rtb_to_xlsx.py"),
                 sid.removeprefix("rtb:"), str(out)], check=True)
        else:
            curl("https://docs.google.com/spreadsheets/d/"
                 f"{sid}/export?format=xlsx", out)
            print(f"downloaded {name}")
    fix_truncated_cubes(src_dir)


def fix_truncated_cubes(src_dir):
    """Maelstrom Pulse's sheet has a truncated Cube tab. Re-saving a
    downloaded sheet with openpyxl drops cached formula values, so copy
    Draft/Matches VALUES into a fresh workbook and synthesize its Cube
    tab from read-the-bones' per-draft card list."""
    import openpyxl
    fixes = {"maelstrom-pulse-(s3)": "maelstrom-pulse"}
    for stem, slug in fixes.items():
        path = src_dir / f"{stem}.xlsx"
        src = openpyxl.load_workbook(path, data_only=True)
        cws = src["Cube"]
        n = 0
        while cws.cell(2 + n, 2).value not in (None, ""):
            n += 1
        if n > 100:
            continue
        cards_path = HERE / "rtb" / f"{slug}-cards.json"
        if not cards_path.exists():
            curl("https://read-the-bones.vercel.app/api/cards?drafts="
                 f"{slug}&poolAsOfDraft={slug}", cards_path)
        cards = json.loads(cards_path.read_text())["cards"]
        wb = openpyxl.Workbook()
        for tab in ("Draft", "Matches"):
            dst = wb.active if tab == "Draft" else wb.create_sheet(tab)
            dst.title = tab
            for row in src[tab].iter_rows():
                for cell in row:
                    if cell.value is not None:
                        dst.cell(cell.row, cell.column, cell.value)
        cws = wb.create_sheet("Cube")
        cws.append(["", "Card", "Type", "Color"])
        for c in cards:
            color = "".join(c.get("colors") or [])
            tline = (c.get("scryfall") or {}).get("typeLine", "")
            cws.append(["", c["cardName"], tline, color])
            for copy in range(2, c.get("maxCopiesInDraft", 1) + 1):
                cws.append(["", f'{c["cardName"]} {copy}', tline, color])
        wb.save(path)
        print(f"{stem}: rebuilt values-only with RTB cube ({len(cards)})")


def main():
    do_verify = "--verify" in sys.argv
    dry_run = "--dry-run" in sys.argv

    src_dir = HERE / "sources"
    download_sources(src_dir)

    out_dir = HERE / "out"
    out_dir.mkdir(exist_ok=True)
    inputs = [f"{n}={src_dir / (slugify(n) + '.xlsx')}@{dpa}"
              for n, (_, dpa) in SOURCES_ALL.items()]
    decks_tsv = HERE / "decks_all.tsv"
    build_cmd = [sys.executable, str(HERE / "roto_summary.py"),
                 "--teams-frac", str(TEAMS_FRAC)]
    if decks_tsv.exists():
        build_cmd += ["--decks", str(decks_tsv)]

    formulas_x = out_dir / "samp-roto-all.xlsx"
    subprocess.run(build_cmd + [str(formulas_x)] + inputs, check=True)

    if do_verify:
        values_x = out_dir / "samp-roto-all-values.xlsx"
        subprocess.run(build_cmd + ["--values", str(values_x)] + inputs,
                       check=True)
        sys.path.insert(0, str(HERE))
        verify(formulas_x, values_x)

    if dry_run:
        print("dry run: skipping publish")
        return
    if not TARGET_SHEET_ID:
        sys.exit("TARGET_SHEET_ID unset: create the sheet once with "
                 "upload_sheet.py, then paste its id into refresh_all.py")
    subprocess.run(
        [sys.executable, str(HERE / "upload_sheet.py"), str(formulas_x),
         TITLE, "--update", TARGET_SHEET_ID], check=True)


if __name__ == "__main__":
    main()
