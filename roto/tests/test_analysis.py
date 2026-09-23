import itertools
import json
import pathlib
import tempfile
import unittest

import openpyxl

from roto.colors import card_colors, is_land
from roto.decks import load_decks
from roto.draft import overall_pick, parse_draft
from roto.mining import Dataset
from roto.packages import (backbone_depth, flex_packages, jaccard_components,
                           package_edges, signature_groups)


class DraftTests(unittest.TestCase):
    def test_pick_order_is_a_per_draft_input(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = pathlib.Path(tmp) / 'draft.xlsx'
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = 'Draft'
            for seat, name in enumerate(['Alice', 'Bob', 'Carol'], 3):
                ws.cell(3, seat, name)
            for rnd in range(1, 7):
                ws.cell(rnd + 3, 1, rnd)
                for seat in range(1, 4):
                    ws.cell(rnd + 3, seat + 2, f'Card {rnd}-{seat}')
            wb.create_sheet('Cube').append(['', 'Card', 'Type', 'Color'])
            wb['Cube'].append(['', 'Card 1-1', 'Creature', 'G'])
            wb.create_sheet('Matches').append(['', 'Matches'])
            wb.save(path)
            snake, _ = parse_draft(path, 'Plain')
            double, _ = parse_draft(path, 'Doubled', dpa=2)
        self.assertEqual(snake.picks['Card 3-1'].overall, 7)
        self.assertEqual(double.picks['Card 3-1'].overall, 7)
        self.assertEqual(snake.picks['Card 4-1'].overall, 12)
        self.assertEqual(double.picks['Card 4-1'].overall, 8)
        self.assertEqual({p.overall for p in double.picks.values()}, set(range(1, 19)))
        self.assertEqual(snake.records, {'Alice': [0, 0], 'Bob': [0, 0], 'Carol': [0, 0]})

    def test_double_pick_traversals_cover_every_position(self):
        for n in [3, 9, 10]:
            for boundary in [20, 23, 24, 25]:
                # Include complete double traversals after the boundary.
                rounds = boundary + 8
                actual = [overall_pick(r, seat, n, boundary)
                          for r in range(1, rounds + 1) for seat in range(1, n + 1)]
                self.assertEqual(sorted(actual), list(range(1, rounds * n + 1)))


class PackageTests(unittest.TestCase):
    def test_full_and_partial_signatures_scale_to_any_pod_count(self):
        sig = ('Alice', 'Bob') + (None,) * 11
        owners = {c: sig for c in ['A', 'B', 'C']}
        self.assertEqual(signature_groups(owners), [])
        self.assertEqual(signature_groups(owners, min_shared=2), [(sig, ['A', 'B', 'C'])])

    def test_shared_absence_does_not_create_an_edge(self):
        groups = [(('Alice', None, None), ['A']),
                  (('Bob', None, None), ['B']),
                  (('Alice', 'Carol', None), ['C'])]
        self.assertEqual(package_edges(groups, agreement=2), [])
        self.assertEqual(package_edges(groups, agreement=1), [(0, 2, [1])])

    def test_flex_tracks_each_live_draft(self):
        owners = {'core': ('A', 'B', 'C', None), 'flex': ('A', 'B', None, None)}
        groups = [(owners['core'], ['core'])]
        result = flex_packages(groups, owners)
        self.assertEqual(result[0]['flex'], [[], [], ['flex'], []])

    def test_jaccard_and_backbone_use_supplied_thresholds(self):
        owners = {'A': ('one', 'two', 'three'), 'B': ('one', 'two', 'three'),
                  'C': ('one', None, None)}
        self.assertEqual(jaccard_components(owners, min_support=2, min_size=2), [{'A', 'B'}])
        self.assertEqual(backbone_depth(owners, set(owners), min_support=2), {'A': 2, 'B': 2, 'C': 0})

    def test_itemsets_match_independent_brute_force(self):
        owners = {'A': ('x', 'x', 'y'), 'B': ('x', 'x', 'z'),
                  'C': ('x', 'z', 'y'), 'D': ('y', 'x', None)}
        dataset = Dataset.from_owners(owners)
        decks = [{p: {c for c, sig in owners.items() if sig[k] == p}
                  for p in {sig[k] for sig in owners.values()} - {None}}
                 for k in range(3)]
        for miss in [0, 1]:
            expected = {}
            for size in range(2, len(owners) + 1):
                for cards in itertools.combinations(owners, size):
                    group = frozenset(cards)
                    support = sum(any(len(group & deck) >= size - miss for deck in pod.values())
                                  for pod in decks)
                    if support >= 2:
                        expected[group] = support
            actual = dict(dataset.mine(min_support=2, min_size=2, miss=miss,
                                       maximal=False, workers=1))
            self.assertEqual(actual, expected)


class MetadataTests(unittest.TestCase):
    def test_color_policy_is_explicit(self):
        metadata = {'Spell': {'colors': ['U'], 'type_line': 'Sorcery',
                              'oracle_text': 'Flashback {R}'},
                    'Rainbow': {'colors': [], 'type_line': 'Land',
                                'produced_mana': list('WUBRG')},
                    'MDFC': {'colors': ['G'], 'type_line': 'Sorcery // Land'}}
        self.assertEqual(card_colors('Spell', metadata), {'U'})
        self.assertEqual(card_colors('Spell', metadata, include_flashback=True), {'U', 'R'})
        self.assertEqual(card_colors('Rainbow', metadata), set())
        self.assertFalse(is_land('MDFC', metadata))

    def test_cached_rebuild_and_cube_alias(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = pathlib.Path(tmp)
            (root / 'deckcache').mkdir()
            (root / 'deckcache/first.json').write_text(json.dumps({'deck': [{'name': 'Other', 'count': 1}]}))
            (root / 'deckcache/rebuild.json').write_text(json.dumps({'deck': [{'name': 'Alias', 'count': 1}]}))
            path = root / 'decks.tsv'
            path.write_text('draft\tplayer\tkind\turl\nPod\tAlice\trebuild\thttps://example/rebuild\n'
                            'Pod\tAlice\tinitial\thttps://example/first\n')
            decks, sizes, links = load_decks(path, [('Canonical', '', ''), ('Other', '', '')],
                                            aliases={'Alias': 'Canonical'})
        self.assertEqual(decks, {('Pod', 'Alice'): {'Canonical': 'main'}})
        self.assertEqual(sizes, {('Pod', 'Alice'): 1})
        self.assertEqual([row[-1] for row in links], ['Y', ''])


if __name__ == '__main__':
    unittest.main()
