"""One-command refresh of the LoL roto Google Sheet.

Downloads fresh exports of the three draft spreadsheets (match results live
in their Matches tabs), fetches any deck pools missing from deckcache/,
rebuilds the workbook, and updates the Google Sheet in place — same URL,
same sharing settings.

  new deck link posted -> add a row to decks.tsv, then: python3 refresh.py
  new matches played   -> python3 refresh.py

Flags:
  --verify   before publishing, also build a --values version, upload the
             formula build to a temporary Sheet, export Google's evaluated
             values, and diff every cell of the computed tabs
  --dry-run  build (and verify) but skip publishing
"""

import pathlib
import subprocess
import sys

HERE = pathlib.Path(__file__).parent
SOURCES = {
    "draft1": "1i5IK8JKOeZpKZpV27YbqIFVNEkOQZq1rXRJQq-IBwCw",
    "draft2": "13ffofLIEuHaFgvHKmB308NNZWQi1YgeVFEKsTlCUc04",
    "draft3": "1AY8wUfcQLltu4D-rFOSfRfcTxuFNgL68EGY26WmaiDU",
}
TARGET_SHEET_ID = "1_w-YcYynXZgzObp13fPUB1q8XNxN6IH7gFkyHgH8E8w"
TITLE = "LoL Cube Roto — Pick Summary"
COMPUTED_TABS = ["Pick Summary", "Win Rates", "Color Analysis",
                 "Deck Shells (all 3 drafts)", "Packages (2 of 3)"]
FUZZY_HEADERS = ("Avg", "Win Rate", "Score")  # float columns: rounding wiggle


def curl(url, out):
    subprocess.run(["curl", "-sL", "--fail", url, "-o", str(out)], check=True)


def download_sources(src_dir):
    src_dir.mkdir(exist_ok=True)
    for name, sid in SOURCES.items():
        curl(f"https://docs.google.com/spreadsheets/d/{sid}/export?format=xlsx",
             src_dir / f"{name}.xlsx")
        print(f"downloaded {name}")


def prefetch_pools():
    """Fetch pools with curl up front (urllib SSL is broken on this python)."""
    cache = HERE / "deckcache"
    cache.mkdir(exist_ok=True)
    for line in (HERE / "decks.tsv").read_text().splitlines()[1:]:
        if not line.strip():
            continue
        pool_id = line.split("\t")[3].rstrip("/").split("/")[-1]
        path = cache / f"{pool_id}.json"
        if not path.exists():
            curl(f"https://sealeddeck.tech/api/pools/{pool_id}", path)
            print(f"fetched pool {pool_id}")


def verify(formulas_path, values_path):
    """Upload formula build to a temp Sheet, export Google's evaluated
    values, and diff against the values build cell by cell."""
    import io

    import openpyxl
    import upload_sheet
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaIoBaseDownload

    drive = build("drive", "v3", credentials=upload_sheet.get_credentials())
    tmp = upload_sheet.upload(str(formulas_path), "tmp-roto-verify")
    try:
        req = drive.files().export_media(
            fileId=tmp["id"],
            mimeType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        buf = io.BytesIO()
        dl = MediaIoBaseDownload(buf, req)
        done = False
        while not done:
            _, done = dl.next_chunk()
    finally:
        drive.files().delete(fileId=tmp["id"]).execute()

    wb_want = openpyxl.load_workbook(values_path, data_only=True)
    wb_got = openpyxl.load_workbook(io.BytesIO(buf.getvalue()), data_only=True)
    total, mismatches = 0, []
    for tab in COMPUTED_TABS:
        want, got = wb_want[tab], wb_got[tab]
        assert want.max_row == got.max_row, (tab, want.max_row, got.max_row)
        headers = [c.value or "" for c in want[1]]
        fuzzy = {i + 1 for i, h in enumerate(headers)
                 if any(k in str(h) for k in FUZZY_HEADERS)}
        for r in range(1, want.max_row + 1):
            for c in range(1, len(headers) + 1):
                total += 1
                w, g = want.cell(r, c).value, got.cell(r, c).value
                if g == "":
                    g = None
                if w == g:
                    continue
                if (c in fuzzy and isinstance(w, (int, float))
                        and isinstance(g, (int, float)) and abs(w - g) <= 0.051):
                    continue
                mismatches.append((tab, r, c, w, g))
    if mismatches:
        for m in mismatches[:20]:
            print("MISMATCH %s row %d col %d: values=%r evaluated=%r" % m)
        sys.exit(f"verification FAILED: {len(mismatches)}/{total} cells differ")
    print(f"verification passed: {total} cells match")


def main():
    do_verify = "--verify" in sys.argv
    dry_run = "--dry-run" in sys.argv

    src_dir = HERE / "sources"
    download_sources(src_dir)
    prefetch_pools()

    out_dir = HERE / "out"
    out_dir.mkdir(exist_ok=True)
    inputs = [str(src_dir / f"{n}.xlsx") for n in SOURCES]
    build_cmd = [sys.executable, str(HERE / "roto_summary.py")]

    formulas_x = out_dir / "lol-roto-summary.xlsx"
    subprocess.run(build_cmd + [str(formulas_x)] + inputs, check=True)

    if do_verify:
        values_x = out_dir / "lol-roto-summary-values.xlsx"
        subprocess.run(build_cmd + ["--values", str(values_x)] + inputs, check=True)
        sys.path.insert(0, str(HERE))
        verify(formulas_x, values_x)

    if dry_run:
        print("dry run: skipping publish")
        return
    subprocess.run(
        [sys.executable, str(HERE / "upload_sheet.py"), str(formulas_x),
         TITLE, "--update", TARGET_SHEET_ID],
        check=True,
    )


if __name__ == "__main__":
    main()
