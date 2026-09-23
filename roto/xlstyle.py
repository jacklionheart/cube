"""Shared spreadsheet palette and cell styling."""
from openpyxl.styles import Alignment, Font, PatternFill

COLOR_FILLS = {
    "W": "F5F4ED", "U": "C0E9F2", "B": "D9D9D9", "R": "F7D7D7",
    "G": "DAFBE0", "C": "EDD7BD", "multi": "F6EAC7",
}
HEADER_FILL = PatternFill("solid", fgColor="434343")
HEADER_FONT = Font(bold=True, color="FFFFFF")
CARD_FONT = Font(bold=True)
CENTER = Alignment(horizontal="center")


def color_fill(color):
    color = (color or "").strip()
    key = color if color in COLOR_FILLS else ("multi" if len(color) > 1 else None)
    return PatternFill("solid", fgColor=COLOR_FILLS[key]) if key else None


def style_card_cell(cell, color):
    fill = color_fill(color)
    if fill:
        cell.fill = fill
    cell.font = CARD_FONT


def style_header(ws):
    for c in ws[1]:
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)
