"""Convert Read the Bones draft data into LoL-template-shaped xlsx.

Two s4 pods (kishla-skimmer, raven-eagle) ran on read-the-bones.vercel.app
instead of the Google Sheet template. This synthesizes the three tabs the
pipeline parses — Draft (snake grid), Cube (card list with the template's
"Name 2" convention for duplicated lands), Matches (round-robin rows) —
from cached API payloads in rtb/:

  <slug>.json            /api/drafts/<slug>/live   (picks, seatNames)
  <slug>-standings.json  /api/drafts/<slug>/standings (matches)
  <slug>-cards.json      /api/cards?drafts=<slug>&poolAsOfDraft=<slug>
                         (cube list; fetched on demand via curl)

Usage: python3 rtb_to_xlsx.py <slug> <out.xlsx>
"""

import json
import pathlib
import subprocess
import sys

import openpyxl

HERE = pathlib.Path(__file__).parent
RTB = HERE / "rtb"
BASE = "https://read-the-bones.vercel.app/api"


def load(slug):
    live = json.loads((RTB / f"{slug}.json").read_text())
    standings = json.loads((RTB / f"{slug}-standings.json").read_text())
    cards_path = RTB / f"{slug}-cards.json"
    if not cards_path.exists():
        url = f"{BASE}/cards?drafts={slug}&poolAsOfDraft={slug}"
        subprocess.run(["curl", "-s", "--fail", url, "-o", str(cards_path)],
                       check=True)
    cards = json.loads(cards_path.read_text())
    return live, standings, cards


def dedupe_names(picks):
    """Second pick of a duplicated card becomes 'Name 2' (template style)."""
    seen, out = set(), []
    for p in sorted(picks, key=lambda p: p["pickN"]):
        name = p["cardName"]
        if name in seen:
            name = f"{name} 2"
        seen.add(name)
        out.append({**p, "cardName": name})
    return out


def convert(slug, out_path):
    live, standings, cards = load(slug)
    n = live["numSeats"]
    rounds = live["picksPerPlayer"]
    players = [live["seatNames"][str(s)] for s in range(1, n + 1)]
    picks = dedupe_names(live["picks"])
    assert len(picks) == n * rounds, (len(picks), n, rounds)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Draft"
    for i, p in enumerate(players):
        ws.cell(3, 3 + i, p)
    per_seat_round = {s: 0 for s in range(1, n + 1)}
    for p in picks:
        per_seat_round[p["seat"]] += 1
        rnd = per_seat_round[p["seat"]]
        ws.cell(3 + rnd, 1, rnd)
        ws.cell(3 + rnd, 2 + p["seat"], p["cardName"])
    assert set(per_seat_round.values()) == {rounds}

    cws = wb.create_sheet("Cube")
    cws.append(["", "Card", "Type", "Color"])
    for c in cards["cards"]:
        color = "".join(c.get("colors") or [])
        tline = c["scryfall"]["typeLine"]
        cws.append(["", c["cardName"], tline, color])
        if c.get("maxCopiesInDraft", 1) > 1:
            cws.append(["", f'{c["cardName"]} 2', tline, color])

    mws = wb.create_sheet("Matches")
    mws.append(["", "Round Robin Tournament"])
    mws.append(["", "Games", "Games Wins"])
    names = {int(k): v for k, v in live["seatNames"].items()}
    for m in standings["matches"]:
        g1, g2 = m["seat1Wins"], m["seat2Wins"]
        mws.append(["", names[m["seat1"]], g1, "VS", g2,
                    names[m["seat2"]], int(g1 > g2), int(g2 > g1)])

    wb.save(out_path)
    print(f"{slug}: {n}p x {rounds}r, cube "
          f"{sum(c.get('maxCopiesInDraft', 1) for c in cards['cards'])}, "
          f"{len(standings['matches'])} matches -> {out_path}")


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
